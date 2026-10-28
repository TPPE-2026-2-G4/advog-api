import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from app.cliente.model import Cliente
from app.processo.model import Processo, StatusProcesso

ROTA = "/processos/exportar"


@pytest.fixture
def auth(token_admin):
    return {"Authorization": f"Bearer {token_admin}"}


@pytest.fixture
def processos(db_session):
    cliente = Cliente(nome="Maria da Silva", email="maria@teste.com", telefone="61999999999")
    db_session.add(cliente)
    db_session.flush()
    itens = [
        Processo(
            descricao="Descrição",
            data_inicio=datetime(2026, 1, 1),
            funcionario_id=1,
            cnj=f"000000000000000000{indice:02d}",
            titulo=titulo,
            tribunal="TJDFT",
            area="Civil",
            status=status,
            data_prazo=datetime(2026, 10, 30),
            cliente_id=cliente.cliente_id,
        )
        for indice, (titulo, status) in enumerate(
            [
                ("Alfa", StatusProcesso.ATIVO),
                ("Beta", StatusProcesso.ARQUIVADO),
                ("Gama", StatusProcesso.ATIVO),
            ]
        )
    ]
    db_session.add_all(itens)
    db_session.commit()
    return itens


def _planilha(response):
    return load_workbook(io.BytesIO(response.content)).active


@pytest.mark.parametrize(
    ("formato", "media_type", "inicio"),
    [
        ("csv", "text/csv", b"\xef\xbb\xbf"),
        ("xlsx", "application/vnd.openxmlformats", b"PK"),
        ("pdf", "application/pdf", b"%PDF"),
    ],
)
def test_exportar_cada_formato(client, auth, processos, formato, media_type, inicio):
    response = client.get(ROTA, params={"formato": formato}, headers=auth)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(media_type)
    assert response.content.startswith(inicio)
    assert "processos_" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith(f'.{formato}"')


def test_exportar_aplica_filtros_e_ignora_paginacao(client, auth, processos):
    todos = client.get(ROTA, params={"formato": "xlsx", "page_size": 1}, headers=auth)
    ativos = client.get(ROTA, params={"formato": "xlsx", "status": "Ativo"}, headers=auth)

    assert _planilha(todos).max_row == 4
    planilha = _planilha(ativos)
    assert [linha[1].value for linha in planilha.iter_rows(min_row=2)] == ["Gama", "Alfa"]


def test_exportar_inclui_nome_do_cliente(client, auth, processos):
    planilha = _planilha(client.get(ROTA, params={"formato": "xlsx"}, headers=auth))

    cabecalho = [c.value for c in planilha[1]]
    assert planilha.cell(2, cabecalho.index("Cliente") + 1).value == "Maria da Silva"
    assert "cliente_id" not in cabecalho


def test_exportar_sem_resultados_gera_somente_cabecalho(client, auth, processos):
    response = client.get(ROTA, params={"formato": "xlsx", "tribunal": "inexistente"}, headers=auth)

    assert response.status_code == 200
    assert _planilha(response).max_row == 1


def test_exportar_exige_formato_valido(client, auth):
    assert client.get(ROTA, headers=auth).status_code == 422
    assert client.get(ROTA, params={"formato": "docx"}, headers=auth).status_code == 422


def test_exportar_respeita_limite_de_linhas(client, auth, processos, monkeypatch):
    monkeypatch.setattr("app.processo.service.LIMITE_LINHAS", 2)

    response = client.get(ROTA, params={"formato": "csv"}, headers=auth)

    assert response.status_code == 422
    assert "refine os filtros" in response.json()["detail"]


def test_exportar_exige_visualizar_processos(client, token_acesso, funcionario_sem_permissao):
    usuario = funcionario_sem_permissao
    token = token_acesso(usuario.funcionario_id, usuario.email)

    response = client.get(
        ROTA, params={"formato": "csv"}, headers={"Authorization": f"Bearer {token}"}
    )

    assert response.status_code == 403


def test_exportar_exige_autenticacao(client):
    assert client.get(ROTA, params={"formato": "csv"}).status_code == 401
