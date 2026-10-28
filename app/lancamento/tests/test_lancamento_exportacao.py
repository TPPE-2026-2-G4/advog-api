import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from app.categoria_lancamento.model import CategoriaLancamento
from app.cliente.model import Cliente
from app.lancamento.model import Lancamento, StatusLancamento, TipoLancamento

ROTA = "/lancamentos/exportar"


@pytest.fixture
def auth(token_admin):
    return {"Authorization": f"Bearer {token_admin}"}


@pytest.fixture
def lancamentos(db_session):
    categoria = CategoriaLancamento(nome_categoria="Honorários")
    cliente = Cliente(nome="Maria da Silva", email="maria@teste.com")
    db_session.add_all([categoria, cliente])
    db_session.flush()
    itens = [
        Lancamento(
            titulo="Honorário A",
            tipo=TipoLancamento.ENTRADA,
            status=StatusLancamento.REALIZADO,
            data_vencimento=datetime(2099, 1, 10),
            data_pagamento=datetime(2099, 1, 9),
            valor=1000.0,
            categoria_id=categoria.categoria_id,
            cliente_id=cliente.cliente_id,
        ),
        Lancamento(
            titulo="Aluguel",
            tipo=TipoLancamento.SAIDA,
            status=StatusLancamento.PENDENTE,
            data_vencimento=datetime(2099, 1, 20),
            valor=400.5,
        ),
    ]
    db_session.add_all(itens)
    db_session.commit()
    return itens


def _planilha(response):
    return load_workbook(io.BytesIO(response.content)).active


@pytest.mark.parametrize(
    ("formato", "inicio"), [("csv", b"\xef\xbb\xbf"), ("xlsx", b"PK"), ("pdf", b"%PDF")]
)
def test_exportar_cada_formato(client, auth, lancamentos, formato, inicio):
    response = client.get(ROTA, params={"formato": formato}, headers=auth)

    assert response.status_code == 200
    assert response.content.startswith(inicio)
    assert "lancamentos_" in response.headers["content-disposition"]


def test_exportar_conteudo_e_totais(client, auth, lancamentos):
    planilha = _planilha(client.get(ROTA, params={"formato": "xlsx"}, headers=auth))

    cabecalho = [c.value for c in planilha[1]]
    linha = {nome: planilha.cell(2, i + 1).value for i, nome in enumerate(cabecalho)}
    assert linha["Título"] == "Honorário A"
    assert linha["Tipo"] == "Entrada"
    assert linha["Status"] == "Realizado"
    assert linha["Categoria"] == "Honorários"
    assert linha["Cliente"] == "Maria da Silva"
    assert linha["Valor"] == 1000.0
    assert planilha.cell(3, cabecalho.index("Tipo") + 1).value == "Saída"
    totais = {
        planilha.cell(i, 1).value: planilha.cell(i, 2).value for i in range(5, planilha.max_row + 1)
    }
    assert totais == {
        "Total de entradas": 1000.0,
        "Total de saídas": 400.5,
        "Saldo": 599.5,
    }


def test_exportar_aplica_filtros_da_listagem(client, auth, lancamentos):
    planilha = _planilha(
        client.get(ROTA, params={"formato": "xlsx", "tipo": "s", "valor_min": 100}, headers=auth)
    )

    assert planilha.cell(2, 1).value == "Aluguel"
    assert planilha.cell(3, 1).value is None
    assert planilha.cell(4, 2).value == 0.0  # total de entradas do conjunto filtrado
    assert planilha.cell(5, 2).value == 400.5


@pytest.mark.parametrize(
    "filtros",
    [{"inicio": "2099-02-01", "fim": "2099-01-01"}, {"valor_min": 10, "valor_max": 5}],
)
def test_exportar_filtros_invalidos(client, auth, filtros):
    response = client.get(ROTA, params={"formato": "csv", **filtros}, headers=auth)

    assert response.status_code == 422


def test_exportar_respeita_limite_de_linhas(client, auth, lancamentos, monkeypatch):
    monkeypatch.setattr("app.lancamento.service.LIMITE_LINHAS", 1)

    response = client.get(ROTA, params={"formato": "csv"}, headers=auth)

    assert response.status_code == 422
    assert "refine os filtros" in response.json()["detail"]


def test_exportar_exige_visualizar_financeiro(client, token_acesso, funcionario_sem_permissao):
    usuario = funcionario_sem_permissao
    token = token_acesso(usuario.funcionario_id, usuario.email)

    response = client.get(
        ROTA, params={"formato": "csv"}, headers={"Authorization": f"Bearer {token}"}
    )

    assert response.status_code == 403


def test_exportar_exige_autenticacao(client):
    assert client.get(ROTA, params={"formato": "csv"}).status_code == 401
