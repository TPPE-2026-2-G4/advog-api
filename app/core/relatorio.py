import csv
import enum
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any
from xml.sax.saxutils import escape

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

LIMITE_LINHAS = 10_000
_PREFIXOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")


class FormatoRelatorio(enum.StrEnum):
    CSV = "csv"
    XLSX = "xlsx"
    PDF = "pdf"


class TipoColuna(enum.StrEnum):
    TEXTO = "texto"
    DATA = "data"
    DATA_HORA = "data_hora"
    MOEDA = "moeda"


@dataclass(frozen=True)
class Coluna:
    titulo: str
    tipo: TipoColuna = TipoColuna.TEXTO


@dataclass
class Tabela:
    titulo: str
    colunas: list[Coluna]
    linhas: list[list[Any]]
    resumo: list[tuple[str, Any]] = field(default_factory=list)


@dataclass(frozen=True)
class Arquivo:
    conteudo: bytes
    media_type: str
    extensao: str


def _protegido(valor: str) -> str:
    return f"'{valor}" if valor.startswith(_PREFIXOS_FORMULA) else valor


def _moeda(valor: float) -> str:
    inteiro, centavos = f"{valor:,.2f}".split(".")
    return f"{inteiro.replace(',', '.')},{centavos}"


def _texto(valor: Any, tipo: TipoColuna) -> str:
    if valor is None:
        return ""
    if tipo is TipoColuna.DATA:
        return valor.strftime("%d/%m/%Y")
    if tipo is TipoColuna.DATA_HORA:
        return valor.strftime("%d/%m/%Y %H:%M")
    if tipo is TipoColuna.MOEDA:
        return _moeda(valor)
    return str(valor)


def _texto_resumo(valor: Any) -> str:
    return _texto(valor, TipoColuna.MOEDA if isinstance(valor, float) else TipoColuna.TEXTO)


def _larguras_pdf(tabela: Tabela, largura_total: float) -> list[float]:
    pesos = []
    for indice, coluna in enumerate(tabela.colunas):
        maior = len(coluna.titulo)
        for linha in tabela.linhas:
            maior = max(maior, len(_texto(linha[indice], coluna.tipo)))
        pesos.append(min(max(maior, 6), 40))
    total = sum(pesos)
    return [largura_total * peso / total for peso in pesos]


def _gerar_csv(tabela: Tabela) -> bytes:
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";", lineterminator="\r\n")
    escritor.writerow([_protegido(c.titulo) for c in tabela.colunas])
    for linha in tabela.linhas:
        escritor.writerow(
            [
                _protegido(texto) if coluna.tipo is TipoColuna.TEXTO else texto
                for coluna, valor in zip(tabela.colunas, linha, strict=True)
                for texto in [_texto(valor, coluna.tipo)]
            ]
        )
    if tabela.resumo:
        escritor.writerow([])
        for rotulo, valor in tabela.resumo:
            escritor.writerow([_protegido(rotulo), _texto_resumo(valor)])
    return b"\xef\xbb\xbf" + saida.getvalue().encode("utf-8")


def _gerar_xlsx(tabela: Tabela) -> bytes:
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = tabela.titulo[:31]
    planilha.append([c.titulo for c in tabela.colunas])
    for celula in planilha[1]:
        celula.font = Font(bold=True)
    planilha.freeze_panes = "A2"

    for linha in tabela.linhas:
        planilha.append(
            [
                _protegido(valor)
                if coluna.tipo is TipoColuna.TEXTO and isinstance(valor, str)
                else valor
                for coluna, valor in zip(tabela.colunas, linha, strict=True)
            ]
        )

    formatos = {
        TipoColuna.DATA: "DD/MM/YYYY",
        TipoColuna.DATA_HORA: "DD/MM/YYYY HH:MM",
        TipoColuna.MOEDA: '"R$" #,##0.00',
    }
    for indice, coluna in enumerate(tabela.colunas, start=1):
        letra = get_column_letter(indice)
        largura = len(coluna.titulo)
        for (celula,) in planilha.iter_rows(min_row=2, min_col=indice, max_col=indice):
            if coluna.tipo in formatos:
                celula.number_format = formatos[coluna.tipo]
            largura = max(largura, len(_texto(celula.value, coluna.tipo)))
        planilha.column_dimensions[letra].width = min(largura + 2, 60)

    if tabela.resumo:
        planilha.append([])
        for rotulo, valor in tabela.resumo:
            planilha.append([rotulo, valor])
            planilha.cell(planilha.max_row, 1).font = Font(bold=True)
            if isinstance(valor, float):
                planilha.cell(planilha.max_row, 2).number_format = formatos[TipoColuna.MOEDA]

    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def _gerar_pdf(tabela: Tabela) -> bytes:
    estilos = getSampleStyleSheet()
    celula = ParagraphStyle(
        "celula", parent=estilos["BodyText"], fontSize=7, leading=9, wordWrap="CJK"
    )
    cabecalho = ParagraphStyle(
        "cabecalho", parent=celula, fontName="Helvetica-Bold", textColor=colors.white
    )

    def paragrafo(texto: str, estilo: ParagraphStyle) -> Paragraph:
        return Paragraph(escape(texto), estilo)

    dados = [[paragrafo(c.titulo, cabecalho) for c in tabela.colunas]]
    for linha in tabela.linhas:
        dados.append(
            [
                paragrafo(_texto(valor, coluna.tipo), celula)
                for coluna, valor in zip(tabela.colunas, linha, strict=True)
            ]
        )

    estilo_tabela = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]
    )
    margem = 1 * cm
    elementos: list[Any] = [
        Paragraph(escape(tabela.titulo), estilos["Title"]),
        Paragraph(
            f"Gerado em {datetime.now():%d/%m/%Y %H:%M} · {len(tabela.linhas)} registro(s)",
            estilos["Normal"],
        ),
        Spacer(1, 0.4 * cm),
        Table(
            dados,
            colWidths=_larguras_pdf(tabela, landscape(A4)[0] - 2 * margem),
            repeatRows=1,
            style=estilo_tabela,
        ),
    ]
    if tabela.resumo:
        elementos.append(Spacer(1, 0.4 * cm))
        for rotulo, valor in tabela.resumo:
            texto = _texto_resumo(valor)
            elementos.append(
                Paragraph(f"<b>{escape(rotulo)}:</b> {escape(texto)}", estilos["Normal"])
            )

    saida = io.BytesIO()
    SimpleDocTemplate(
        saida,
        pagesize=landscape(A4),
        leftMargin=margem,
        rightMargin=margem,
        topMargin=margem,
        bottomMargin=margem,
        title=tabela.titulo,
    ).build(elementos)
    return saida.getvalue()


_GERADORES = {
    FormatoRelatorio.CSV: (_gerar_csv, "text/csv; charset=utf-8"),
    FormatoRelatorio.XLSX: (
        _gerar_xlsx,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ),
    FormatoRelatorio.PDF: (_gerar_pdf, "application/pdf"),
}


def gerar_arquivo(tabela: Tabela, formato: FormatoRelatorio) -> Arquivo:
    gerador, media_type = _GERADORES[formato]
    return Arquivo(conteudo=gerador(tabela), media_type=media_type, extensao=formato.value)


def resposta_arquivo(arquivo: Arquivo, nome_base: str) -> Response:
    nome = f"{nome_base}_{date.today():%Y%m%d}.{arquivo.extensao}"
    return Response(
        content=arquivo.conteudo,
        media_type=arquivo.media_type,
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )
